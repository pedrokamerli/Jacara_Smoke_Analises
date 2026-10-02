"""Cadastro mínimo privado: nome + HMAC. Contatos não são persistidos."""
import io
import pandas as pd
from openpyxl import load_workbook
from .privacy import pseudonymize_phone
from .profile_sales import _normalize_header


def prepare_customer_directory(payload,secret):
    book=load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
    names={}
    try:
        sheet=book[book.sheetnames[0]]
        sheet.reset_dimensions()
        rows=sheet.iter_rows(values_only=True)
        columns={_normalize_header(value):i for i,value in enumerate(next(rows,()))}
        if not {'nome','telefone principal'}<=columns.keys():
            raise ValueError('Cadastro precisa de Nome e Telefone Principal')
        for row in rows:
            get=lambda field:row[columns[field]] if columns[field]<len(row) else None
            key=pseudonymize_phone(get('telefone principal'),source='pos',secret=secret)
            name=' '.join(str(get('nome') or '').split())
            if key and name:
                names.setdefault(key,set()).add(name)
    finally: book.close()
    records=[{'customer_key':key,'customer_name':next(iter(values))} for key,values in names.items() if len(values)==1]
    return pd.DataFrame(records,columns=['customer_key','customer_name']),{'available':True,'matched_identifiers':len(records),'ambiguous_identifiers_omitted':sum(len(values)>1 for values in names.values())}


def render_customer_ranking(sq,manifest,root,table):
    import streamlit as st
    artifact=manifest.get('customer_directory')
    if not artifact:
        st.info('Para mostrar nomes no privado, envie também Lista-Clientes. Os indicadores agregados continuam disponíveis.')
        return
    path=(root/artifact).resolve()
    if not path.is_relative_to((root/'data/runs').resolve()) or path.is_symlink():
        raise ValueError('Cadastro privado fora do diretório autorizado')
    names=pd.read_parquet(path)
    if set(names.columns)!={'customer_key','customer_name'} or names.customer_key.duplicated().any():
        raise ValueError('Contrato do cadastro privado inválido')
    frame=sq("select customer_key,count(*) as orders,sum(amount_received_brl) as received,min(cast(opened_at as date)) as first_purchase,max(cast(opened_at as date)) as last_purchase from selected_orders where customer_key is not null group by 1")
    frame=frame.merge(names,on='customer_key',how='inner',validate='one_to_one')
    st.subheader('Melhores clientes e recorrência no recorte')
    st.caption('Cadastro confidencial, disponível somente nesta área autenticada. Melhor por valor = maior valor recebido acumulado, não lucro. Recorrente = duas ou mais compras no período selecionado. Telefones, endereços e códigos protegidos não são exibidos.')
    if frame.empty:
        st.info('Nenhum cliente do cadastro foi ligado aos pedidos selecionados.')
        return
    cols=st.columns(3)
    rank=cols[0].selectbox('Classificar clientes por',['Valor acumulado','Quantidade de compras'],key='named_customer_rank')
    recurring=cols[1].checkbox('Somente recorrentes',key='named_recurring_only')
    search=cols[2].text_input('Buscar cliente por nome',key='named_customer_search')
    if recurring: frame=frame.loc[frame.orders>=2]
    if search.strip(): frame=frame.loc[frame.customer_name.str.contains(search.strip(),case=False,regex=False)]
    order=['received','orders'] if rank=='Valor acumulado' else ['orders','received']
    frame=frame.sort_values(order+['customer_name'],ascending=[False,False,True])
    frame['ticket']=frame.received/frame.orders
    frame['recurring']=frame.orders.ge(2).map({True:'Sim',False:'Não'})
    display=frame.drop(columns='customer_key').rename(columns={'customer_name':'Cliente','orders':'Compras no recorte','received':'Valor acumulado (R$)','ticket':'Ticket médio (R$)','first_purchase':'Primeira compra no recorte','last_purchase':'Última compra no recorte','recurring':'Recorrente'})
    table(display[['Cliente','Compras no recorte','Valor acumulado (R$)','Ticket médio (R$)','Recorrente','Primeira compra no recorte','Última compra no recorte']].head(50))
    st.caption(f'{len(frame)} clientes encontrados; exibindo até 50. Cadastro com nomes ambíguos para um mesmo contato é omitido. Os filtros globais também controlam este ranking; primeira compra é a primeira dentro do recorte, não da vida do cliente.')
