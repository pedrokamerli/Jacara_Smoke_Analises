"""Perfilamento local agregado das planilhas principais de vendas.

Não grava chaves de pedidos, nomes, contatos, endereços ou texto livre no
relatório. O relatório gerado permanece em data/interim (ignorado pelo Git).
"""

from __future__ import annotations

import argparse
import io
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

DEFAULT_ARCHIVE = Path("JACARE_SMOKE_HOUSE_PEDIDOS_E_RELATORIOS_ORGANIZADOS_2026.zip")
ORDERS_MEMBER = "01_Todos_os_pedidos_01-01_a_20-08-2026.xlsx"
ITEMS_MEMBER = "02_Historico_Itens_Vendidos_01-01_a_20-08-2026.xlsx"
HISTORY_MEMBER = "04_PEDIDOS_HISTORICOS.xlsx"


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().casefold())


def _normalize_order_key(value: Any) -> str | None:
    if value is None:
        return None
    key = str(value).strip()
    if not key:
        return None
    if re.fullmatch(r"\d+\.0+", key):
        key = key.split(".", 1)[0]
    return key.casefold()


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            parsed = from_excel(value)
            return parsed.date() if isinstance(parsed, datetime) else parsed
        except (OverflowError, TypeError, ValueError):
            return None
    if isinstance(value, str):
        raw = value.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
    return None


def _as_number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        raw = value.strip().replace("R$", "").replace(" ", "")
        if not raw:
            return None
        if "," in raw:
            raw = raw.replace(".", "").replace(",", ".")
        try:
            return float(raw)
        except ValueError:
            return None
    return None


def _find_member(archive: zipfile.ZipFile, suffix: str) -> str:
    matches = [name for name in archive.namelist() if name.endswith(suffix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"Esperava encontrar uma planilha para `{suffix}`; encontrei {len(matches)}."
        )
    return matches[0]


def _workbook_rows(payload: bytes):
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, ())
    columns = {_normalize_header(value): index for index, value in enumerate(headers)}
    return workbook, rows, columns


def _required(columns: dict[str, int], *names: str) -> dict[str, int]:
    missing = [name for name in names if _normalize_header(name) not in columns]
    if missing:
        raise RuntimeError("Colunas necessárias não encontradas: " + ", ".join(missing))
    return {_normalize_header(name): columns[_normalize_header(name)] for name in names}


def _get(row: tuple[Any, ...], columns: dict[str, int], name: str) -> Any:
    index = columns[_normalize_header(name)]
    return row[index] if index < len(row) else None


def _read_orders(payload: bytes, closed_cutoff: date) -> dict[str, Any]:
    workbook, rows, columns = _workbook_rows(payload)
    wanted = _required(columns, "Código", "Data Abertura", "Status", "Origem", "Total Recebido")
    keys: set[str] = set()
    status_counts: Counter[str] = Counter()
    origin_counts: Counter[str] = Counter()
    dates: list[date] = []
    duplicate_rows = missing_key_rows = paid_rows_closed = rows_after_cutoff = 0
    paid_total_closed = 0.0

    try:
        for row in rows:
            key = _normalize_order_key(_get(row, wanted, "Código"))
            if key is None:
                missing_key_rows += 1
            elif key in keys:
                duplicate_rows += 1
            else:
                keys.add(key)

            status_value = _get(row, wanted, "Status")
            status = re.sub(r"\s+", " ", str(status_value or "").strip())
            if status:
                status_counts[status] += 1
            origin_value = _get(row, wanted, "Origem")
            origin = re.sub(r"\s+", " ", str(origin_value or "").strip())
            if origin:
                origin_counts[origin] += 1
            opened = _as_date(_get(row, wanted, "Data Abertura"))
            if opened:
                dates.append(opened)
                if opened > closed_cutoff:
                    rows_after_cutoff += 1
            if status.casefold() == "finalizado - pago":
                if opened is not None and opened <= closed_cutoff:
                    paid_rows_closed += 1
                    amount = _as_number(_get(row, wanted, "Total Recebido"))
                    if amount is not None:
                        paid_total_closed += amount
    finally:
        workbook.close()

    return {
        "keys": keys,
        "rows_with_key": len(keys) + duplicate_rows,
        "unique_order_keys": len(keys),
        "duplicate_order_rows": duplicate_rows,
        "missing_order_key_rows": missing_key_rows,
        "date_min": min(dates) if dates else None,
        "date_max": max(dates) if dates else None,
        "status_counts": status_counts,
        "origin_counts": origin_counts,
        "closed_cutoff": closed_cutoff,
        "rows_after_cutoff": rows_after_cutoff,
        "paid_rows_through_cutoff": paid_rows_closed,
        "paid_total_received_through_cutoff": paid_total_closed,
    }


def _read_order_keys(payload: bytes, column_name: str) -> tuple[set[str], int, int]:
    workbook, rows, columns = _workbook_rows(payload)
    wanted = _required(columns, column_name)
    keys: set[str] = set()
    duplicate_rows = missing_rows = 0
    try:
        for row in rows:
            key = _normalize_order_key(_get(row, wanted, column_name))
            if key is None:
                missing_rows += 1
            elif key in keys:
                duplicate_rows += 1
            else:
                keys.add(key)
    finally:
        workbook.close()
    return keys, duplicate_rows, missing_rows


def _read_items(payload: bytes) -> dict[str, Any]:
    workbook, rows, columns = _workbook_rows(payload)
    wanted = _required(columns, "Cod. Ped.", "Qtd.", "Stat. Ped.")
    order_keys: set[str] = set()
    status_counts: Counter[str] = Counter()
    item_rows = missing_order_key_rows = 0
    quantity_sum = 0.0
    try:
        for row in rows:
            item_rows += 1
            key = _normalize_order_key(_get(row, wanted, "Cod. Ped."))
            if key is None:
                missing_order_key_rows += 1
            else:
                order_keys.add(key)
            status_value = _get(row, wanted, "Stat. Ped.")
            status = re.sub(r"\s+", " ", str(status_value or "").strip())
            if status:
                status_counts[status] += 1
            quantity = _as_number(_get(row, wanted, "Qtd."))
            if quantity is not None:
                quantity_sum += quantity
    finally:
        workbook.close()

    return {
        "item_rows": item_rows,
        "order_keys": order_keys,
        "missing_order_key_rows": missing_order_key_rows,
        "status_counts": status_counts,
        "quantity_sum": quantity_sum,
    }


def profile_archive(archive_path: Path, closed_cutoff: date) -> str:
    with zipfile.ZipFile(archive_path) as outer:
        orders_path = _find_member(outer, ORDERS_MEMBER)
        items_path = _find_member(outer, ITEMS_MEMBER)
        history_path = _find_member(outer, HISTORY_MEMBER)
        orders = _read_orders(outer.read(orders_path), closed_cutoff)
        items = _read_items(outer.read(items_path))
        historical_keys, historical_duplicates, historical_missing = _read_order_keys(
            outer.read(history_path), "Código"
        )

    current_keys: set[str] = orders["keys"]
    item_keys: set[str] = items["order_keys"]
    overlap = current_keys & historical_keys
    matched_item_keys = current_keys & item_keys
    unmatched_item_keys = item_keys - current_keys
    current_only = current_keys - historical_keys
    history_only = historical_keys - current_keys

    lines = [
        "# Perfilamento local agregado de vendas",
        "",
        "> Gerado localmente. Contém somente contagens, categorias operacionais e agregados; não contém registros, chaves de pedido ou dados pessoais.",
        "",
        "## Pedidos consolidados",
        "",
        f"- Linhas com código de pedido: {orders['rows_with_key']:,}",
        f"- Códigos únicos: {orders['unique_order_keys']:,}",
        f"- Linhas repetindo código: {orders['duplicate_order_rows']:,}",
        f"- Linhas sem código: {orders['missing_order_key_rows']:,}",
        f"- Intervalo de abertura: {_fmt_date(orders['date_min'])} a {_fmt_date(orders['date_max'])}",
        f"- Corte adotado para comparação: {orders['closed_cutoff'].isoformat()}",
        f"- Registros abertos após o corte (período parcial separado): {orders['rows_after_cutoff']:,}",
        f"- Pedidos pagos abertos até o corte: {orders['paid_rows_through_cutoff']:,}",
        f"- Soma `Total Recebido` desses pedidos: R$ {orders['paid_total_received_through_cutoff']:,.2f}",
        "",
        "### Contagem por status",
        "",
        *_counter_lines(orders["status_counts"]),
        "",
        "### Contagem por origem registrada",
        "",
        *_counter_lines(orders["origin_counts"]),
        "",
        "## Itens vendidos",
        "",
        f"- Linhas de item: {items['item_rows']:,}",
        f"- Soma de quantidades: {items['quantity_sum']:,.2f}",
        f"- Linhas sem código de pedido: {items['missing_order_key_rows']:,}",
        f"- Códigos de pedido distintos nos itens: {len(item_keys):,}",
        f"- Códigos de item presentes na base consolidada: {len(matched_item_keys):,}",
        f"- Códigos de item sem correspondência na base consolidada: {len(unmatched_item_keys):,}",
        "",
        "### Contagem por status do pedido nos itens",
        "",
        *_counter_lines(items["status_counts"]),
        "",
        "## Reconciliação com a base histórica",
        "",
        f"- Códigos únicos na base histórica: {len(historical_keys):,}",
        f"- Linhas repetindo código na base histórica: {historical_duplicates:,}",
        f"- Linhas sem código na base histórica: {historical_missing:,}",
        f"- Códigos presentes nas duas bases: {len(overlap):,}",
        f"- Códigos apenas na base consolidada: {len(current_only):,}",
        f"- Códigos apenas na base histórica: {len(history_only):,}",
        "",
        "## Observações",
        "",
        "- A soma acima usa `Total Recebido` apenas para pedidos `Finalizado - Pago` abertos até o corte; precisa ser comparada com o relatório oficial antes de virar métrica publicada.",
        "- A comparação de itens mede cobertura de códigos, não reconcilia valores nem prova que os itens sejam completos.",
        "- A base histórica pode conter uma janela diferente; as contagens de interseção não devem ser somadas.",
        "- Nenhum campo de cliente, contato, endereço ou observação livre foi incluído no relatório.",
    ]
    return "\n".join(lines) + "\n"


def _fmt_date(value: date | None) -> str:
    return value.isoformat() if value else "não identificado"


def _counter_lines(counter: Counter[str]) -> list[str]:
    if not counter:
        return ["- Sem valores preenchidos."]
    return [f"- `{label}`: {count:,}" for label, count in sorted(counter.items())]


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Perfilamento agregado local de pedidos e itens; não emite dados pessoais."
    )
    parser.add_argument("archive", nargs="?", type=Path, default=DEFAULT_ARCHIVE)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/interim/perfilamento_local.md"),
    )
    parser.add_argument(
        "--cutoff",
        type=date.fromisoformat,
        default=date(2026, 8, 19),
        help="Última data completa no formato AAAA-MM-DD (padrão: 2026-08-19).",
    )
    args = parser.parse_args()
    report = profile_archive(args.archive, args.cutoff)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(report, encoding="utf-8")
    print(f"Perfilamento agregado salvo localmente em {args.output}.")


if __name__ == "__main__":
    main()
