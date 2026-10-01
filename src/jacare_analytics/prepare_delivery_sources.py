"""Prepara agregados de delivery sem persistir IDs, nomes ou avaliações textuais."""

from __future__ import annotations

import argparse
import io
import json
import re
import unicodedata
import zipfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .profile_sales import DEFAULT_ARCHIVE, _as_number, _normalize_header

APPDELIVERY_MEMBER = "Pedidos_AppDelivery_01-01_a_04-08-2026.xlsx"
NINEFOOD_MEMBER = "_Dados do pedido(01-07-2026_31-07-2026).xlsx"


def _safe_category(value: Any) -> str | None:
    if value is None:
        return None
    normalized = unicodedata.normalize("NFKD", str(value).strip().casefold())
    plain = normalized.encode("ascii", "ignore").decode("ascii")
    result = re.sub(r"[^a-z0-9]+", "_", plain).strip("_")
    return result or None


def _as_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)):
        number = int(value)
        if 19_000_101 <= number <= 21_001_231:
            try:
                return datetime.strptime(str(number), "%Y%m%d")
            except ValueError:
                return None
        try:
            parsed = from_excel(value)
            if isinstance(parsed, datetime):
                return parsed
            if isinstance(parsed, date):
                return datetime.combine(parsed, datetime.min.time())
        except (OverflowError, TypeError, ValueError):
            return None
    if isinstance(value, str):
        raw = value.strip()
        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d",
            "%Y%m%d",
            "%d/%m/%Y %H:%M:%S",
            "%d/%m/%Y %H:%M",
            "%d/%m/%Y",
        ):
            try:
                return datetime.strptime(raw, fmt)
            except ValueError:
                continue
    return None


def _duration_number(value: Any) -> float | None:
    if isinstance(value, str):
        raw = value.strip()
        if re.fullmatch(r"\d{1,3}(?:\.\d{3})+", raw):
            return float(raw.replace(".", ""))
    return _as_number(value)


def _read_sheet(payload: bytes):
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, ())
    columns = {_normalize_header(value): index for index, value in enumerate(headers)}
    return workbook, rows, columns


def _required(columns: dict[str, int], *names: str) -> dict[str, int]:
    missing = [name for name in names if _normalize_header(name) not in columns]
    if missing:
        raise RuntimeError("Colunas necessárias ausentes no export de delivery.")
    return {_normalize_header(name): columns[_normalize_header(name)] for name in names}


def _get(row: tuple[Any, ...], columns: dict[str, int], name: str) -> Any:
    index = columns[_normalize_header(name)]
    return row[index] if index < len(row) else None


def _find_member(archive: zipfile.ZipFile, suffix: str) -> str:
    matches = [name for name in archive.namelist() if name.endswith(suffix)]
    if len(matches) != 1:
        raise RuntimeError(f"Esperava um arquivo correspondente a `{suffix}`; encontrei {len(matches)}.")
    return matches[0]


def _prepare_appdelivery(payload: bytes) -> tuple[pd.DataFrame, dict[str, int]]:
    workbook, rows, columns = _read_sheet(payload)
    selected = _required(columns, "Criado em", "Valor Itens", "Valor Entrega", "Tipo de Entrega", "Status")
    daily: dict[tuple[date, str, str], dict[str, float]] = defaultdict(
        lambda: {"orders": 0, "items_value_brl": 0.0, "delivery_fee_brl": 0.0}
    )
    missing_date_rows = 0
    try:
        for row in rows:
            created = _as_datetime(_get(row, selected, "Criado em"))
            if created is None:
                missing_date_rows += 1
                continue
            delivery_type = _safe_category(_get(row, selected, "Tipo de Entrega")) or "unknown"
            status = _safe_category(_get(row, selected, "Status")) or "unknown"
            bucket = daily[(created.date(), delivery_type, status)]
            bucket["orders"] += 1
            for target, field in [("items_value_brl", "Valor Itens"), ("delivery_fee_brl", "Valor Entrega")]:
                value = _as_number(_get(row, selected, field))
                if value is None:
                    raise ValueError("Valor necessário ausente no AppDelivery; não será substituído por zero.")
                bucket[target] += value
    finally:
        workbook.close()

    records = [
        {
            "order_date": day,
            "delivery_type": delivery_type,
            "order_status": status,
            **values,
        }
        for (day, delivery_type, status), values in sorted(daily.items())
    ]
    return pd.DataFrame.from_records(
        records,
        columns=["order_date", "delivery_type", "order_status", "orders", "items_value_brl", "delivery_fee_brl"],
    ), {"appdelivery_rows_aggregated": sum(int(row["orders"]) for row in records), "appdelivery_rows_without_date": missing_date_rows}


def _prepare_99food(payload: bytes) -> tuple[pd.DataFrame, dict[str, int]]:
    workbook, rows, columns = _read_sheet(payload)
    money_fields = {
        "sales_revenue_brl": "Receita de vendas",
        "shop_revenue_brl": "Receita real da loja",
        "offer_expenses_brl": "Despesas de ofertas da loja",
        "commission_expense_brl": "Despesas de comissão da loja",
        "payment_channel_fee_brl": "Taxa de canal de pagamento da loja",
        "platform_rewards_brl": "Recompensas da plataforma",
        "original_delivery_fee_brl": "Taxa de entrega original da loja",
        "new_customer_delivery_fee_brl": "Taxa de entrega paga por novos clientes",
        "refund_brl": "Valor do reembolso",
        "free_delivery_net_cost_brl": "Custo líquido da loja na oferta de entrega grátis",
        "logistics_cost_brl": "Custos logísticos",
    }
    duration_fields = {
        "prep_minutes": "Tempo de preparo (minutos)",
        "acceptance_seconds": "Tempo para aceitação do pedido (segundos)",
        "finalization_seconds": "Tempo de finalização do pedido (segundos)",
        "delivery_seconds": "Duração da entrega (segundos)",
    }
    required = [
        "Data", "Horário da conclusão", "Horário do cancelamento", "Contagem do item",
        "Preparação atrasada?", "Nível de avaliação do cliente",
        *money_fields.values(), *duration_fields.values(),
    ]
    selected = _required(columns, *required)
    daily: dict[date, dict[str, Any]] = defaultdict(
        lambda: {
            "orders": 0,
            "orders_with_completion_timestamp": 0,
            "orders_with_cancellation_timestamp": 0,
            "orders_with_both_timestamps": 0,
            "late_preparation_orders": 0,
            "items_count": 0.0,
            "rating_count": 0,
            "rating_sum": 0.0,
            **{field: 0.0 for field in money_fields},
            **{f"{field}_sum": 0.0 for field in duration_fields},
            **{f"{field}_count": 0 for field in duration_fields},
        }
    )
    missing_date_rows = 0
    try:
        for row in rows:
            order_date_time = _as_datetime(_get(row, selected, "Data"))
            if order_date_time is None:
                missing_date_rows += 1
                continue
            bucket = daily[order_date_time.date()]
            bucket["orders"] += 1
            has_completion = _as_datetime(_get(row, selected, "Horário da conclusão")) is not None
            has_cancellation = _as_datetime(_get(row, selected, "Horário do cancelamento")) is not None
            bucket["orders_with_completion_timestamp"] += has_completion
            bucket["orders_with_cancellation_timestamp"] += has_cancellation
            bucket["orders_with_both_timestamps"] += has_completion and has_cancellation
            bucket["late_preparation_orders"] += _safe_category(_get(row, selected, "Preparação atrasada?")) == "sim"
            item_count = _as_number(_get(row, selected, "Contagem do item"))
            if item_count is None:
                raise ValueError("Quantidade ausente no 99Food; não será substituída por zero.")
            bucket["items_count"] += item_count
            rating = _as_number(_get(row, selected, "Nível de avaliação do cliente"))
            if rating is not None:
                bucket["rating_count"] += 1
                bucket["rating_sum"] += rating
            for output_field, source_field in money_fields.items():
                amount = _as_number(_get(row, selected, source_field))
                if amount is None:
                    raise ValueError("Valor necessário ausente no 99Food; não será substituído por zero.")
                bucket[output_field] += amount
            for output_field, source_field in duration_fields.items():
                duration = _duration_number(_get(row, selected, source_field))
                if duration is not None and duration > 0:
                    bucket[f"{output_field}_sum"] += duration
                    bucket[f"{output_field}_count"] += 1
    finally:
        workbook.close()

    records = []
    for order_date, values in sorted(daily.items()):
        record = {"order_date": order_date}
        for key, value in values.items():
            record[key] = value
        record["average_rating"] = (
            values["rating_sum"] / values["rating_count"]
            if values["rating_count"] >= 5
            else None
        )
        for field in duration_fields:
            denominator = values[f"{field}_count"]
            record[f"average_{field}"] = (
                values[f"{field}_sum"] / denominator if denominator >= 5 else None
            )
        records.append(record)

    frame = pd.DataFrame.from_records(records)
    metadata = {
        "99food_rows_aggregated": sum(int(row["orders"]) for row in records),
        "99food_days_aggregated": len(records),
        "99food_rows_without_date": missing_date_rows,
    }
    return frame, metadata


def prepare_delivery_archive(archive_path: Path, output_dir: Path) -> dict[str, int]:
    with zipfile.ZipFile(archive_path) as archive:
        app_member = _find_member(archive, APPDELIVERY_MEMBER)
        food_member = _find_member(archive, NINEFOOD_MEMBER)
        appdelivery, app_metadata = _prepare_appdelivery(archive.read(app_member))
        food99, food_metadata = _prepare_99food(archive.read(food_member))

    output_dir.mkdir(parents=True, exist_ok=True)
    app_path = output_dir / "appdelivery_daily.parquet"
    food_path = output_dir / "food99_daily.parquet"
    appdelivery.to_parquet(app_path, index=False, engine="pyarrow")
    food99.to_parquet(food_path, index=False, engine="pyarrow")
    summary = {
        "source_archive": archive_path.name,
        "appdelivery_source_member": Path(app_member).name,
        "food99_source_member": Path(food_member).name,
        "appdelivery_parquet": str(app_path),
        "food99_parquet": str(food_path),
        **app_metadata,
        **food_metadata,
    }
    summary_path = Path("data/interim/delivery_prepare_summary.json")
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return {key: value for key, value in summary.items() if isinstance(value, int)}


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera agregados locais de delivery sem salvar registros individuais.")
    parser.add_argument("archive", nargs="?", type=Path, default=DEFAULT_ARCHIVE)
    parser.add_argument("--output-dir", type=Path, default=Path("data/processed"))
    args = parser.parse_args()
    summary = prepare_delivery_archive(args.archive, args.output_dir)
    print(
        "Agregados de delivery salvos localmente: "
        f"{summary.get('appdelivery_rows_aggregated', 0)} pedidos AppDelivery e "
        f"{summary.get('99food_rows_aggregated', 0)} pedidos 99Food."
    )


if __name__ == "__main__":
    main()
