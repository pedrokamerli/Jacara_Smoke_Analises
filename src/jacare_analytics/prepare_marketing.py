"""Extrai apenas métricas aditivas do detalhamento diário do Meta Ads."""

from __future__ import annotations

import io
from collections import defaultdict
from datetime import date
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from .profile_sales import _as_date, _as_number, _normalize_header


def prepare_meta_ads(payload: bytes, cutoff: date) -> tuple[pd.DataFrame, dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    daily: dict[date, dict[str, float]] = defaultdict(lambda: {"spend_brl": 0.0, "impressions": 0.0, "link_clicks": 0.0, "link_clicks_observed_rows": 0, "link_clicks_missing_rows": 0})
    summary_spend = None
    source_detail_spend = 0.0
    daily_rows = skipped_summary_rows = 0
    try:
        if "Raw Data Report" not in workbook.sheetnames:
            raise ValueError("A fonte Meta Ads precisa da aba Raw Data Report.")
        rows = workbook["Raw Data Report"].iter_rows(values_only=True)
        columns = None
        for row in rows:
            headers = {_normalize_header(value): index for index, value in enumerate(row)}
            if {"dia", "valor gasto (brl)", "nível de veiculação", "impressões", "cliques no link"}.issubset(headers):
                columns = headers
                break
        if columns is None:
            raise ValueError("Cabeçalho do Meta Ads não corresponde ao contrato de dados atual.")
        level = None
        for row in rows:
            def get(field: str):
                index = columns[field]
                return row[index] if index < len(row) else None

            row_level = str(get("nível de veiculação") or "").strip().lower()
            if row_level:
                level = row_level
            day = _as_date(get("dia"))
            spend = _as_number(get("valor gasto (brl)"))
            if day is None:
                skipped_summary_rows += 1
                if level is None and spend is not None:
                    if summary_spend is not None:
                        raise ValueError("Mais de um total geral encontrado no Meta Ads.")
                    summary_spend = spend
                continue
            if level != "ad":
                raise ValueError("Detalhamento diário do Meta Ads apareceu fora do nível de anúncio.")
            if spend is None:
                raise ValueError("Há gasto ausente numa linha diária do Meta Ads.")
            source_detail_spend += spend
            if day > cutoff:
                continue
            impressions = _as_number(get("impressões"))
            if impressions is None:
                raise ValueError("Há impressões ausentes numa linha diária do Meta Ads; não serão substituídas por zero.")
            clicks = _as_number(get("cliques no link"))
            daily[day]["spend_brl"] += spend
            daily[day]["impressions"] += impressions
            daily[day]["link_clicks_missing_rows"] += clicks is None
            if clicks is not None:
                daily[day]["link_clicks"] += clicks
                daily[day]["link_clicks_observed_rows"] += 1
            daily_rows += 1
    finally:
        workbook.close()
    frame = pd.DataFrame.from_records([
        {"metric_date": day, **values, "link_clicks": values["link_clicks"] if values["link_clicks_observed_rows"] else None} for day, values in sorted(daily.items())
    ], columns=["metric_date", "spend_brl", "impressions", "link_clicks", "link_clicks_observed_rows", "link_clicks_missing_rows"])
    frame = frame.astype({"metric_date":"datetime64[ns]", "spend_brl":float, "impressions":float, "link_clicks":float, "link_clicks_observed_rows":int, "link_clicks_missing_rows":int})
    total_spend = round(float(frame["spend_brl"].sum()), 2)
    return frame, {
        "daily_ad_rows": daily_rows,
        "summary_rows_excluded": skipped_summary_rows,
        "daily_spend_brl": total_spend,
        "report_total_spend_brl": summary_spend,
        "source_detail_spend_brl": round(source_detail_spend, 2),
        "spend_reconciled": summary_spend is not None and abs(round(source_detail_spend, 2) - summary_spend) < 0.01,
        "missing_link_click_rows": int(frame["link_clicks_missing_rows"].sum()),
        "link_click_policy": "Somente cliques observados são somados; valores ausentes permanecem desconhecidos. CPC e CTR só são calculados com cobertura completa.",
        "reach_policy": "Alcance não foi somado porque pode repetir pessoas entre anúncios e datas.",
    }
