"""Transforma planilhas brutas em Parquet minimizado para DuckDB/dbt.

Os ZIPs são lidos diretamente; nenhum arquivo bruto é extraído. A saída segue
allowlists explícitas, remove identificadores diretos e é gravada em data/,
que é local e está ignorada pelo Git.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .profile_sales import (
    DEFAULT_ARCHIVE,
    ITEMS_MEMBER,
    ORDERS_MEMBER,
    _as_number,
    _find_member,
    _normalize_header,
    _normalize_order_key,
    _required,
)
from .privacy import KEY_ENV_VAR, pseudonymize_phone

DEFAULT_CUTOFF = date(2026, 8, 19)


def _read_sheet(payload: bytes) -> tuple[Any, Iterator[tuple[Any, ...]], dict[str, int]]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, ())
    columns = {_normalize_header(value): index for index, value in enumerate(headers)}
    return workbook, rows, columns


def _excel_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)):
        try:
            result = from_excel(value)
            if isinstance(result, datetime):
                return result
            if isinstance(result, date):
                return datetime.combine(result, datetime.min.time())
        except (OverflowError, TypeError, ValueError):
            return None
    if isinstance(value, str):
        raw = value.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y"):
            try:
                return datetime.strptime(raw, fmt)
            except ValueError:
                continue
    return None


def _cell(row: tuple[Any, ...], columns: dict[str, int], field: str) -> Any:
    index = columns[_normalize_header(field)]
    return row[index] if index < len(row) else None


def _clean_category(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    if not text:
        return None
    normalized = text.casefold()
    if normalized in {"finalizado - pago", "finalizado pago"}:
        return "paid"
    if normalized in {"em andamento", "em aberto"}:
        return "open"
    if normalized in {"finalizado - fiado", "finalizado fiado"}:
        return "credit"
    ascii_text = unicodedata.normalize("NFKD", normalized).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", ascii_text).strip("_") or None


def _clean_label(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _build_orders(
    payload: bytes, cutoff: date, hmac_secret: str | None
) -> tuple[pd.DataFrame, dict[str, int], dict[str, int]]:
    workbook, rows, columns = _read_sheet(payload)
    needed = [
        "Código", "Data Abertura", "Status", "Tipo", "Origem", "Tipo Entrega",
        "Data Fechamento", "Tot. Itens", "Serviço", "Valor Entrega", "Total",
        "Total Recebido", "Telefone",
    ]
    selected = _required(columns, *needed)
    records: list[dict[str, Any]] = []
    order_key_map: dict[str, int] = {}
    seen: set[str] = set()
    duplicate_rows = missing_key_rows = invalid_date_rows = after_cutoff_rows = 0
    phone_candidate_rows = valid_customer_keys = 0

    try:
        for row in rows:
            source_code = _normalize_order_key(_cell(row, selected, "Código"))
            if not source_code:
                missing_key_rows += 1
                continue
            if source_code in seen:
                duplicate_rows += 1
                continue
            seen.add(source_code)

            opened_at = _excel_datetime(_cell(row, selected, "Data Abertura"))
            if opened_at is None:
                invalid_date_rows += 1
                continue
            if opened_at.date() > cutoff:
                after_cutoff_rows += 1
                continue

            order_key = len(order_key_map) + 1
            order_key_map[source_code] = order_key
            raw_phone = _cell(row, selected, "Telefone")
            customer_key = None
            if raw_phone is not None and str(raw_phone).strip():
                phone_candidate_rows += 1
                if hmac_secret:
                    customer_key = pseudonymize_phone(raw_phone, source="pos", secret=hmac_secret)
                    valid_customer_keys += customer_key is not None

            records.append(
                {
                    "order_key": order_key,
                    "opened_at": opened_at,
                    "closed_at": _excel_datetime(_cell(row, selected, "Data Fechamento")),
                    "order_status": _clean_category(_cell(row, selected, "Status")),
                    "order_type": _clean_category(_cell(row, selected, "Tipo")),
                    "source_channel": _clean_category(_cell(row, selected, "Origem")),
                    "delivery_type": _clean_category(_cell(row, selected, "Tipo Entrega")),
                    "items_amount_brl": _as_number(_cell(row, selected, "Tot. Itens")),
                    "service_fee_brl": _as_number(_cell(row, selected, "Serviço")),
                    "delivery_fee_brl": _as_number(_cell(row, selected, "Valor Entrega")),
                    "order_total_brl": _as_number(_cell(row, selected, "Total")),
                    "amount_received_brl": _as_number(_cell(row, selected, "Total Recebido")),
                    "customer_key": customer_key,
                }
            )
    finally:
        workbook.close()

    frame = pd.DataFrame.from_records(
        records,
        columns=[
            "order_key", "opened_at", "closed_at", "order_status", "order_type",
            "source_channel", "delivery_type", "items_amount_brl", "service_fee_brl",
            "delivery_fee_brl", "order_total_brl", "amount_received_brl", "customer_key",
        ],
    )
    metadata = {
        "orders_written": len(frame),
        "duplicate_order_rows_skipped": duplicate_rows,
        "missing_order_key_rows_skipped": missing_key_rows,
        "invalid_order_date_rows_skipped": invalid_date_rows,
        "rows_after_cutoff_skipped": after_cutoff_rows,
        "rows_with_phone_candidate": phone_candidate_rows,
        "rows_with_valid_pseudonym": valid_customer_keys,
    }
    return frame, metadata, order_key_map


def _build_items(
    payload: bytes, order_key_map: dict[str, int], cutoff: date
) -> tuple[pd.DataFrame, dict[str, int]]:
    workbook, rows, columns = _read_sheet(payload)
    needed = [
        "Data/Hora Item", "Qtd.", "Valor Un. Item", "Valor. Tot. Item", "Tipo de Item",
        "Nome Prod", "Tipo Prod", "Cat. Prod.", "Cod. Ped.", "Data Ab. Ped.",
        "Data Fec. Ped.", "Tipo Ped.", "Stat. Ped.",
    ]
    selected = _required(columns, *needed)
    records: list[dict[str, Any]] = []
    missing_order_key_rows = unmatched_order_rows = invalid_date_rows = after_cutoff_rows = 0
    try:
        for row in rows:
            source_code = _normalize_order_key(_cell(row, selected, "Cod. Ped."))
            if not source_code:
                missing_order_key_rows += 1
                continue
            sold_at = _excel_datetime(_cell(row, selected, "Data/Hora Item"))
            order_opened_at = _excel_datetime(_cell(row, selected, "Data Ab. Ped."))
            if sold_at is None and order_opened_at is None:
                invalid_date_rows += 1
                continue
            event_date = sold_at or order_opened_at
            if event_date.date() > cutoff:
                after_cutoff_rows += 1
                continue
            order_key = order_key_map.get(source_code)
            if order_key is None:
                unmatched_order_rows += 1
                continue

            records.append(
                {
                    "order_key": order_key,
                    "sold_at": sold_at or order_opened_at,
                    "quantity": _as_number(_cell(row, selected, "Qtd.")),
                    "unit_price_brl": _as_number(_cell(row, selected, "Valor Un. Item")),
                    "line_total_brl": _as_number(_cell(row, selected, "Valor. Tot. Item")),
                    "item_type": _clean_category(_cell(row, selected, "Tipo de Item")),
                    "product_name": _clean_label(_cell(row, selected, "Nome Prod")),
                    "product_type": _clean_category(_cell(row, selected, "Tipo Prod")),
                    "product_category": _clean_label(_cell(row, selected, "Cat. Prod.")),
                    "order_type": _clean_category(_cell(row, selected, "Tipo Ped.")),
                    "order_status": _clean_category(_cell(row, selected, "Stat. Ped.")),
                }
            )
    finally:
        workbook.close()

    frame = pd.DataFrame.from_records(
        records,
        columns=[
            "order_key", "sold_at", "quantity", "unit_price_brl", "line_total_brl",
            "item_type", "product_name", "product_type", "product_category",
            "order_type", "order_status",
        ],
    )
    metadata = {
        "items_written": len(frame),
        "item_rows_without_order_code": missing_order_key_rows,
        "item_rows_without_order_match": unmatched_order_rows,
        "item_rows_without_date": invalid_date_rows,
        "item_rows_after_cutoff": after_cutoff_rows,
    }
    return frame, metadata


def prepare_archive(archive_path: Path, cutoff: date, output_dir: Path) -> dict[str, int]:
    hmac_secret = os.environ.get(KEY_ENV_VAR)
    if hmac_secret and len(hmac_secret.encode("utf-8")) < 32:
        raise ValueError(f"{KEY_ENV_VAR} deve ter pelo menos 32 bytes; o valor não será exibido.")

    with zipfile.ZipFile(archive_path) as archive:
        orders_member = _find_member(archive, ORDERS_MEMBER)
        items_member = _find_member(archive, ITEMS_MEMBER)
        orders, order_metadata, order_map = _build_orders(
            archive.read(orders_member), cutoff, hmac_secret
        )
        items, item_metadata = _build_items(archive.read(items_member), order_map, cutoff)

    output_dir.mkdir(parents=True, exist_ok=True)
    orders_path = output_dir / "orders.parquet"
    items_path = output_dir / "items.parquet"
    orders.to_parquet(orders_path, index=False, engine="pyarrow")
    items.to_parquet(items_path, index=False, engine="pyarrow")

    report: dict[str, Any] = {
        "source_archive": archive_path.name,
        "cutoff": cutoff.isoformat(),
        "orders_parquet": str(orders_path),
        "items_parquet": str(items_path),
        "customer_pseudonymization_enabled": bool(hmac_secret),
        **order_metadata,
        **item_metadata,
    }
    summary_dir = Path("data/interim")
    summary_dir.mkdir(parents=True, exist_ok=True)
    (summary_dir / "prepare_sources_summary.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return {key: value for key, value in report.items() if isinstance(value, int)}


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepara tabelas Parquet sem identificadores diretos.")
    parser.add_argument("archive", nargs="?", type=Path, default=DEFAULT_ARCHIVE)
    parser.add_argument("--cutoff", type=date.fromisoformat, default=DEFAULT_CUTOFF)
    parser.add_argument("--output-dir", type=Path, default=Path("data/processed"))
    args = parser.parse_args()
    summary = prepare_archive(args.archive, args.cutoff, args.output_dir)
    print(
        "Tabelas Parquet minimizadas criadas localmente; "
        f"{summary.get('orders_written', 0)} pedidos e {summary.get('items_written', 0)} itens."
    )


if __name__ == "__main__":
    main()
